"""Phase 4 Task 4: the two rich writers - DOCX and the self-contained HTML bundle.

Both are pure functions `write(doc, options) -> bytes` over the same
`TranscriptDoc` the six text writers take (ADR-003: paragraphs are computed
from the words on every call; nothing grouped is stored, and no writer sees
a connection, a file or a request). The document is built by hand from the
seed transcript through the helpers `tests/test_exports_text.py` already
has, so every byte is a function of things pinned in that file. No database,
no GPU.

What a golden cannot explain is asserted directly: that the DOCX opens with
python-docx and has the title as its first heading, one paragraph per
rendered paragraph, a bold run per speaker and a hyperlink element per
timestamp; that the HTML bundle points at nothing on the network, carries
its audio as a ``data:`` URL or a relative sidecar name, keeps every
sentence's ``data-start`` for the inline player, and parses as HTML with
every tag balanced. The DOCX golden is the ``word/document.xml`` member - a
byte golden of the whole zip would break on another zlib - and the bundle's
golden is the file itself. Regenerate with `SCRIBE_UPDATE_GOLDENS=1`.
"""

from __future__ import annotations

import base64
import copy
import io
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from html.parser import HTMLParser

import pytest
from docx import Document
from docx.oxml.ns import qn
from markupsafe import escape

from scribe import exports, render
from scribe.exports import docx as docx_writer
from scribe.exports import html_bundle
from scribe.exports.options import ExportOptions
from seed import default_words
from test_exports_text import NOON, check_golden, doc_of, seed_doc, words_of

# --- helpers -------------------------------------------------------------------


def opened(data: bytes):
    """The DOCX bytes as python-docx reads them back."""
    return Document(io.BytesIO(data))


def hyperlinks_in(paragraph) -> list[tuple[str, str]]:
    """``(address, text)`` of every hyperlink element in the paragraph, in order."""
    return [(link.address, link.text) for link in paragraph.hyperlinks]


def bold_runs(paragraph) -> list[str]:
    return [run.text for run in paragraph.runs if run.bold]


def document_xml(data: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("word/document.xml")


def html_of(data: bytes) -> str:
    return data.decode("utf-8")


class Balance(HTMLParser):
    """Counts tags and checks that every one that opens also closes, in order."""

    VOID = {"meta", "link", "br", "img", "input", "hr", "source", "wbr"}

    def __init__(self) -> None:
        super().__init__()
        self.stack: list[str] = []
        self.counts: Counter[str] = Counter()
        self.errors: list[str] = []
        self.ids: list[str] = []

    def handle_starttag(self, tag, attrs):
        self.counts[tag] += 1
        for name, value in attrs:
            if name == "id" and value:
                self.ids.append(value)
        if tag not in self.VOID:
            self.stack.append(tag)

    def handle_endtag(self, tag):
        if not self.stack or self.stack[-1] != tag:
            self.errors.append(f"</{tag}> closes {self.stack[-1] if self.stack else 'nothing'}")
        else:
            self.stack.pop()


def parsed(text: str) -> Balance:
    parser = Balance()
    parser.feed(text)
    parser.close()
    return parser


# --- the registry ----------------------------------------------------------------


def test_docx_and_html_are_registered_with_their_writers():
    assert exports.FORMATS["docx"].writer is docx_writer.write
    assert exports.FORMATS["html"].writer is html_bundle.write
    assert exports.FORMATS["docx"].extension == "docx"
    assert exports.FORMATS["html"].extension == "html"
    assert all(entry.writer is not NotImplemented for entry in exports.FORMATS.values())


@pytest.mark.parametrize("writer", [docx_writer.write, html_bundle.write])
def test_rich_writers_are_pure_and_deterministic(writer):
    doc = seed_doc()
    snapshot = copy.deepcopy(doc)
    options = ExportOptions(timestamps="sentence")

    first = writer(doc, options)
    second = writer(doc, options)

    assert isinstance(first, bytes)
    assert first == second
    assert doc == snapshot


# --- DOCX ------------------------------------------------------------------------


def test_docx_opens_and_starts_with_the_title_as_a_heading():
    document = opened(docx_writer.write(seed_doc(), ExportOptions()))

    first = document.paragraphs[0]
    assert first.text == "Clip"
    assert first.style.name.startswith("Heading")
    assert document.core_properties.title == "Clip"


def test_docx_has_the_facts_line_under_the_title():
    document = opened(docx_writer.write(seed_doc(), ExportOptions()))

    assert document.paragraphs[1].text == "Clip.wav · 0:30 · large-v3-turbo · en · 2026-09-02"


def test_docx_has_one_paragraph_per_rendered_paragraph_with_a_bold_speaker_run():
    doc = seed_doc()
    rendered = render.paragraphs(doc.words, doc.labels)

    body = opened(docx_writer.write(doc, ExportOptions())).paragraphs[2:]

    assert len(body) == len(rendered) == 2
    assert [b.strip() for p in body for b in bold_runs(p)] == ["Arthur:", "Speaker 2:"]
    assert body[0].text == (
        "0:00 Arthur: Don't panic, the towel is still the most important item."
        " The answer to life, the universe and everything is forty-two."
    )
    assert body[1].text == (
        "0:10 Speaker 2: Marvin says the improbability drive makes him even more depressed."
        " Vogon poetry is the third worst in the known universe."
    )


def test_docx_timestamps_are_hyperlink_elements_to_the_scribe_deep_link():
    data = docx_writer.write(seed_doc(), ExportOptions())

    body = opened(data).paragraphs[2:]
    assert [hyperlinks_in(p) for p in body] == [
        [("scribe://media/1#t=0", "0:00")],
        [("scribe://media/1#t=10", "0:10")],
    ]
    # The element itself, not only python-docx's reading of it.
    assert [len(p._p.findall(qn("w:hyperlink"))) for p in body] == [1, 1]
    assert docx_writer.link_for(seed_doc(), 12.345) == "scribe://media/1#t=12.345"


def test_docx_sentence_timestamps_link_every_sentence():
    body = opened(docx_writer.write(seed_doc(), ExportOptions(timestamps="sentence"))).paragraphs[2:]

    assert hyperlinks_in(body[0]) == [("scribe://media/1#t=0", "0:00"), ("scribe://media/1#t=5", "0:05")]
    assert hyperlinks_in(body[1]) == [("scribe://media/1#t=10", "0:10"), ("scribe://media/1#t=15", "0:15")]
    assert body[0].text.startswith("0:00 Arthur: Don't panic, the towel is still the most important item. 0:05 The answer")


def test_docx_word_timestamps_link_every_word():
    body = opened(docx_writer.write(seed_doc(), ExportOptions(timestamps="word"))).paragraphs[2:]

    assert sum(len(hyperlinks_in(p)) for p in body) == len(default_words())
    assert body[0].text.startswith("0:00 Arthur: Don't 0:00 panic, 0:01 the 0:01 towel 0:02 is")


def test_docx_without_timestamps_has_no_hyperlinks():
    data = docx_writer.write(seed_doc(), ExportOptions(timestamps="none"))

    document = opened(data)
    assert all(not p.hyperlinks for p in document.paragraphs)
    assert b"w:hyperlink" not in document_xml(data)
    assert document.paragraphs[2].text.startswith("Arthur: Don't panic")


def test_docx_without_speakers_has_no_bold_runs():
    document = opened(docx_writer.write(seed_doc(), ExportOptions(speakers=False)))

    body = document.paragraphs[2:]
    assert all(not bold_runs(p) for p in body)
    assert body[0].text.startswith("0:00 Don't panic")
    assert "Arthur" not in "".join(p.text for p in document.paragraphs)


def test_docx_timestamp_format_changes_the_link_text_not_the_target():
    body = opened(docx_writer.write(seed_doc(), ExportOptions(timestamp_format="smpte"))).paragraphs[2:]

    assert hyperlinks_in(body[1]) == [("scribe://media/1#t=10", "00:00:10:00")]


def test_docx_text_survives_as_text():
    words = words_of("AT&T beats <b>bold</b> & \"quotes\".", speaker="SPEAKER_00")
    doc = doc_of(words, labels={"SPEAKER_00": "R&D <lead>"}, title="Notes <1> & 2")

    document = opened(docx_writer.write(doc, ExportOptions()))

    assert document.paragraphs[0].text == "Notes <1> & 2"
    assert document.paragraphs[2].text == "0:00 R&D <lead>: AT&T beats <b>bold</b> & \"quotes\"."


def test_docx_core_properties_name_the_transcript():
    props = opened(docx_writer.write(seed_doc(), ExportOptions())).core_properties

    assert props.author == "MyScribe"
    assert props.last_modified_by == "MyScribe"
    assert props.created == datetime.fromtimestamp(NOON, tz=timezone.utc)
    assert props.modified == props.created
    assert props.language == "en"


def test_docx_bytes_are_deterministic_and_the_document_xml_matches_its_golden():
    data = docx_writer.write(seed_doc(), ExportOptions())

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert archive.testzip() is None
        assert archive.namelist()[0] == "[Content_Types].xml"
        assert {info.date_time for info in archive.infolist()} == {(1980, 1, 1, 0, 0, 0)}
    check_golden("seed.docx.xml", document_xml(data))


def test_docx_ignores_the_text_encoding_and_line_ending_options():
    doc = seed_doc()

    plain = docx_writer.write(doc, ExportOptions())
    other = docx_writer.write(doc, ExportOptions(encoding="cp1252", crlf=True))

    assert plain == other


def test_docx_of_an_empty_transcript_is_the_title_and_the_facts():
    document = opened(docx_writer.write(doc_of([]), ExportOptions()))

    assert [p.text for p in document.paragraphs] == ["Clip", "Clip.wav · 0:30 · large-v3-turbo · en · 2026-09-02"]


def test_docx_of_a_run_without_diarization_has_no_bold_runs():
    document = opened(docx_writer.write(doc_of(words_of("Don't panic. Mostly harmless.")), ExportOptions()))

    body = document.paragraphs[2:]
    assert len(body) == 1
    assert not bold_runs(body[0])
    assert body[0].text == "0:00 Don't panic. Mostly harmless."


# --- the HTML bundle ---------------------------------------------------------------


def test_html_bundle_matches_its_golden():
    data = html_bundle.write(seed_doc(), ExportOptions())

    check_golden("seed.html", data)


def test_html_bundle_references_nothing_on_the_network():
    audio = bytes(range(256)) * 4
    for data in (
        html_bundle.write(seed_doc(), ExportOptions()),
        html_bundle.write(seed_doc(), ExportOptions(), audio=audio),
    ):
        text = html_of(data)
        assert "http://" not in text
        assert "https://" not in text
        assert "<link" not in text
        assert 'src="/' not in text
        assert "htmx" not in text
        assert "hx-" not in text


def test_html_bundle_embeds_audio_as_a_data_url_when_given():
    audio = b"\x00\x01binary\xff"
    expected = base64.b64encode(audio).decode("ascii")

    mp4 = html_of(html_bundle.write(seed_doc(), ExportOptions(), audio=audio))
    mp3 = html_of(html_bundle.write(seed_doc(), ExportOptions(), audio=audio, audio_mime="audio/mpeg"))

    assert f'<audio id="player" controls preload="metadata" src="data:audio/mp4;base64,{expected}"' in mp4
    assert f'src="data:audio/mpeg;base64,{expected}"' in mp3
    assert "Clip.m4a" not in mp4


def test_html_bundle_points_at_a_relative_sidecar_when_no_audio_is_given():
    doc = seed_doc()

    text = html_of(html_bundle.write(doc, ExportOptions()))

    assert 'src="Clip.m4a"' in text
    assert "data:" not in text
    assert html_bundle.sidecar_name(doc, ExportOptions()) == "Clip.m4a"
    assert html_bundle.sidecar_name(doc, ExportOptions(), "audio/mpeg") == "Clip.mp3"
    assert html_bundle.sidecar_name(doc, ExportOptions(), "audio/wav") == "Clip.wav"
    assert html_bundle.sidecar_name(doc, ExportOptions(filename_template="{title}-{lang}"), "audio/ogg") == "Clip-en.ogg"


def test_html_bundle_sidecar_name_is_scrubbed_and_url_quoted():
    doc = doc_of(default_words(), title="My Clip: take #2")

    text = html_of(html_bundle.write(doc, ExportOptions()))

    assert html_bundle.sidecar_name(doc, ExportOptions()) == "My Clip_ take #2.m4a"
    assert 'src="My%20Clip_%20take%20%232.m4a"' in text


@pytest.mark.parametrize(
    ("mime", "extension"),
    [
        ("audio/mp4", "m4a"),
        ("audio/mpeg", "mp3"),
        ("audio/wav", "wav"),
        ("audio/x-wav", "wav"),
        ("audio/aac", "aac"),
        ("audio/ogg", "ogg"),
        ("audio/flac", "flac"),
        ("video/webm", "webm"),
        ("video/mp4", "mp4"),
        ("audio/mp4; codecs=mp4a", "m4a"),
        ("audio/x-something-else", "somethingelse"),
        ("", "bin"),
    ],
)
def test_audio_extension_for_a_mime_type(mime, extension):
    assert html_bundle.audio_extension(mime) == extension


def test_html_bundle_carries_every_sentence_start_and_every_word_span():
    doc = seed_doc()
    sentences = [s for p in render.paragraphs(doc.words, doc.labels) for s in p.sentences]

    text = html_of(html_bundle.write(doc, ExportOptions()))

    stamps = re.findall(r'<a class="ts" href="#t=([^"]+)" data-start="([^"]+)">([^<]*)</a>', text)
    assert [s[1] for s in stamps] == [str(s.start) for s in sentences]
    assert [s[0] for s in stamps] == [str(s.start) for s in sentences]
    assert [s[2] for s in stamps] == ["0:00", "0:05", "0:10", "0:15"]
    spans = re.findall(
        r'<span class="w band-(\w+)" data-i="(\d+)" data-s="([^"]+)" data-e="([^"]+)">([^<]*)</span>', text
    )
    assert len(spans) == len(default_words())
    for word, (band, idx, start, end, spelled) in zip(default_words(), spans):
        assert band == render.confidence_band(word["probability"])
        assert int(idx) == word["idx"]
        assert start == str(word["start"]) and end == str(word["end"])
        # The leading space stays on the word; the text is escaped as text.
        assert spelled == str(escape(word["text"]))


def test_html_bundle_parses_with_balanced_tags_and_one_inline_script():
    text = html_of(html_bundle.write(seed_doc(), ExportOptions(), audio=b"\x00" * 64))

    result = parsed(text)
    assert result.errors == []
    assert result.stack == []
    assert result.counts["script"] == 1
    assert result.counts["style"] == 1
    assert result.counts["audio"] == 1
    assert result.counts["h3"] == 2
    assert result.counts["section"] == 2
    assert {"player", "transcript"} <= set(result.ids)
    assert text.lstrip().lower().startswith("<!doctype html>")
    assert text.count("<script") == 1 and text.count("</script>") == 1
    script = text[text.index("<script") : text.index("</script>")]
    assert "currentTime" in script and "data-start" in script and "timeupdate" in script


def test_html_bundle_has_the_title_the_facts_and_the_language():
    text = html_of(html_bundle.write(seed_doc(), ExportOptions()))

    assert '<html lang="en">' in text
    assert '<meta charset="utf-8">' in text
    assert "<title>Clip</title>" in text
    assert "<h1>Clip</h1>" in text
    assert "Clip.wav · 0:30 · large-v3-turbo · en · 2026-09-02" in text


def test_html_bundle_speaker_headings_only_with_speakers_on():
    doc = seed_doc()

    on = html_of(html_bundle.write(doc, ExportOptions()))
    off = html_of(html_bundle.write(doc, ExportOptions(speakers=False)))

    headings = re.findall(r'<h3 class="speaker" data-cluster="([^"]+)">([^<]*)</h3>', on)
    assert headings == [("SPEAKER_00", "Arthur"), ("SPEAKER_01", "Speaker 2")]
    assert 'data-speaker="SPEAKER_00"' in on
    assert "<h3" not in off
    assert "Arthur" not in off and "Speaker 2" not in off
    assert "SPEAKER_00" not in off


def test_html_bundle_hides_timestamps_when_none_but_keeps_the_seek_targets():
    doc = seed_doc()

    shown = html_of(html_bundle.write(doc, ExportOptions()))
    hidden = html_of(html_bundle.write(doc, ExportOptions(timestamps="none")))

    assert "<body>" in shown
    assert '<body class="hide-ts">' in hidden
    assert hidden.count('data-start="') == shown.count('data-start="') == 4


def test_html_bundle_escapes_text_names_and_the_title():
    words = words_of("AT&T beats <b>bold</b>.", speaker="SPEAKER_00")
    doc = doc_of(words, labels={"SPEAKER_00": "R&D <lead>"}, title="Notes <1> & 2")

    text = html_of(html_bundle.write(doc, ExportOptions()))

    assert "<title>Notes &lt;1&gt; &amp; 2</title>" in text
    assert "<h1>Notes &lt;1&gt; &amp; 2</h1>" in text
    assert ">R&amp;D &lt;lead&gt;</h3>" in text
    assert "&lt;b&gt;bold&lt;/b&gt;" in text
    assert "<b>" not in text
    assert parsed(text).errors == []


def test_html_bundle_is_utf8_without_a_bom_whatever_the_encoding_option():
    words = words_of("Un café, s'il vous plaît.", speaker="SPEAKER_00")
    doc = doc_of(words, labels={"SPEAKER_00": "Zoë"})

    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        data = html_bundle.write(doc, ExportOptions(encoding=encoding))
        assert not data.startswith(b"\xef\xbb\xbf"), encoding
        assert "café" in data.decode("utf-8"), encoding
        assert b"\xc3\xa9" in data, encoding


def test_html_bundle_honours_crlf():
    doc = seed_doc()

    lf = html_bundle.write(doc, ExportOptions())
    crlf = html_bundle.write(doc, ExportOptions(crlf=True))

    assert b"\r" not in lf
    assert crlf.replace(b"\r\n", b"\n") == lf
    assert not crlf.replace(b"\r\n", b"").count(b"\n")


def test_html_bundle_of_an_empty_transcript_still_parses():
    text = html_of(html_bundle.write(doc_of([]), ExportOptions()))

    result = parsed(text)
    assert result.errors == [] and result.stack == []
    assert "<title>Clip</title>" in text
    assert 'data-start="' not in text
    assert "This transcript has no words." in text
    assert result.counts["audio"] == 1


def test_html_bundle_of_a_run_without_diarization_has_no_headings():
    text = html_of(html_bundle.write(doc_of(words_of("Don't panic. Mostly harmless.")), ExportOptions()))

    assert "<h3" not in text
    assert text.count('<section class="para">') == 1
    assert text.count('data-start="') == 2
